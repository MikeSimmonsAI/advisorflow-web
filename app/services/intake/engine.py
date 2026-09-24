"""THE UNIVERSAL INTAKE ENGINE.

    SOURCE CONNECTOR (CSV, Excel, HubSpot export, Google Contacts, API, ...)
        |  rows of {header: value}
        v
    create_batch()          ImportBatch + ImportBatchFile, headers, suggested mapping
    save_mapping()          operator-confirmed mapping, classification, update policy
    run_analysis()          PARSE -> NORMALIZE -> MATCH/DEDUPE -> CLASSIFY
                            -> OUTREACH STATUS -> ANALYZE
        v
    import_staged_rows      one row per source row, nothing in the CRM yet
        v
    commit.py               ONLY on an explicit decision, never automatically

Every source feeds the SAME functions below. A connector's only job is to turn
its source into headers + rows; it never cleans, dedupes or classifies on its
own. That is how the platform avoids a separate, disagreeing cleaning path per
connector.

NOTHING IN THIS MODULE WRITES A LEAD OR AN ORG CONTACT.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Tuple

from sqlalchemy.orm import Session

from app.models.import_models import (ImportBatch, ImportBatchStatus, ImportDuplicateStatus,
                                      ImportMatchConfidence, ImportRowReviewStatus,
                                      ImportStagedRow, ImportValidationStatus)
from app.models.intake_models import (DuplicateResolution, EmailStatus, ImportBatchFile,
                                      IntakeClassification, IntakeStatus, MatchType,
                                      SmsStatus)
from app.services.intake import classification as C
from app.services.intake import eligibility as E
from app.services.intake import fields as F
from app.services.intake import matching as M
from app.services.intake import normalize as N

log = logging.getLogger(__name__)

PIPELINE = "universal"
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_ROWS = 250_000
INSERT_CHUNK = 1000
STALE_AFTER = timedelta(minutes=10)

STAGES = ("uploading", "parsing", "normalizing", "matching", "deduplicating",
          "classifying", "analyzing", "ready_for_review")

SOURCE_OPTIONS = [
    ("hubspot", "HubSpot"), ("csv", "CSV Import"), ("excel", "Excel Import"),
    ("google_contacts", "Google Contacts"), ("salesforce", "Salesforce"),
    ("gohighlevel", "GoHighLevel"), ("website", "Website"), ("referral", "Referral"),
    ("facebook", "Facebook"), ("carrier_file", "Carrier File"),
    ("purchased_list", "Purchased List"), ("manual", "Manual Entry"), ("api", "API"),
    ("lead_scraper", "Lead Scraper"), ("partner", "Partner Integration"), ("other", "Other"),
]
SOURCE_LABELS = dict(SOURCE_OPTIONS)


class IntakeError(ValueError):
    """A problem the operator can fix (bad file, bad mapping)."""


# ══════════════════════════════════════════════════════════════════════════
# 1. READING A SOURCE INTO HEADERS + ROWS
# ══════════════════════════════════════════════════════════════════════════

def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = content.decode(enc)
            if enc == "utf-16" and "\x00" in text[:200]:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return content.decode("cp1252", errors="replace")


def _unique_headers(raw: List[str]) -> List[str]:
    out, seen = [], {}
    for i, h in enumerate(raw):
        h = N.clean_text(h) or f"Column {i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        out.append(h)
    return out


def read_table(content: bytes, filename: str) -> Tuple[List[str], List[List[str]], dict]:
    """Return (headers, rows, stats). Rows are lists aligned to headers; a row
    with MORE cells than headers keeps the extras (reported as malformed)."""
    name = (filename or "").lower()
    stats = {"blank_rows": 0, "malformed_rows": [], "encoding_note": None}
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise IntakeError("Excel support is not installed on this server.") from exc
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise IntakeError(f"This Excel file could not be opened: {str(exc)[:120]}")
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        try:
            header_row = next(it)
        except StopIteration:
            raise IntakeError("The spreadsheet is empty.")
        headers = _unique_headers([("" if v is None else str(v)) for v in header_row])
        rows = []
        for vals in it:
            cells = []
            for v in vals:
                if v is None:
                    cells.append("")
                elif isinstance(v, datetime):
                    cells.append(v.isoformat(sep=" "))
                elif isinstance(v, float) and v.is_integer():
                    cells.append(str(int(v)))
                else:
                    cells.append(str(v))
            if not any(c.strip() for c in cells):
                stats["blank_rows"] += 1
                continue
            rows.append(cells)
        return headers, rows, stats
    if name.endswith(".xls"):
        raise IntakeError("Legacy .xls files are not supported. Save as .xlsx or .csv.")

    text = _decode(content)
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    try:
        headers = _unique_headers(next(reader))
    except StopIteration:
        raise IntakeError("The file is empty.")
    except csv.Error as exc:
        raise IntakeError(f"The header row could not be read: {exc}")
    rows = []
    line = 1
    while True:
        try:
            cells = next(reader)
        except StopIteration:
            break
        except csv.Error as exc:
            line += 1
            stats["malformed_rows"].append({"line": line, "problem": str(exc)[:120]})
            continue
        line += 1
        if not any((c or "").strip() for c in cells):
            stats["blank_rows"] += 1
            continue
        if len(cells) > len(headers):
            extra = cells[len(headers):]
            if any(e.strip() for e in extra):
                stats["malformed_rows"].append({"line": line,
                                                "problem": f"{len(extra)} extra cell(s)"})
        rows.append(cells)
    if len(rows) > MAX_ROWS:
        raise IntakeError(f"This file has {len(rows):,} rows; the limit per batch is "
                          f"{MAX_ROWS:,}. Split it and import in parts.")
    return headers, rows, stats


# ══════════════════════════════════════════════════════════════════════════
# 2. CREATING A BATCH
# ══════════════════════════════════════════════════════════════════════════

def _org_prefix(ctx) -> str:
    import re
    words = re.sub(r"[^A-Za-z0-9 ]", " ", ctx.org_name or ctx.org_slug or "ORG").split()
    p = words[0][:3].upper() if words else "ORG"
    return p if len(p) == 3 else (p + "XX")[:3]


def next_batch_code(db: Session, ctx) -> str:
    """ATL-20260924-001 style: org prefix, date, per-org per-day sequence."""
    day = datetime.utcnow().strftime("%Y%m%d")
    prefix = f"{_org_prefix(ctx)}-{day}-"
    n = (db.query(ImportBatch)
         .filter(ImportBatch.organization_id == ctx.org_id,
                 ImportBatch.batch_code.like(prefix + "%")).count())
    return f"{prefix}{n + 1:03d}"


def org_custom_field_keys(db: Session, org_id: str) -> Dict[str, str]:
    from app.models.models import Organization
    org = db.query(Organization).filter(Organization.id == org_id).first()
    out = {}
    try:
        defs = json.loads(org.crm_custom_fields) if org and org.crm_custom_fields else []
    except Exception:  # noqa: BLE001
        defs = []
    for d in defs if isinstance(defs, list) else []:
        if isinstance(d, dict) and d.get("key"):
            out[F.compact(d["key"])] = d["key"]
            if d.get("label"):
                out[F.compact(d["label"])] = d["key"]
    return out


def org_catalog(db: Session, org_id: str):
    rows = (db.query(IntakeClassification)
            .filter(IntakeClassification.organization_id == org_id).all())
    return C.catalog(rows)


def create_batch(db: Session, ctx, *, content: bytes, filename: str, source: str,
                 source_detail: Optional[str] = None, source_year: Optional[int] = None,
                 list_name: Optional[str] = None, campaign_purpose: Optional[str] = None,
                 offer_hook: Optional[str] = None, tags: Optional[List[str]] = None,
                 display_name: Optional[str] = None) -> ImportBatch:
    if not content:
        raise IntakeError("The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise IntakeError("File too large. Maximum upload size is 50 MB.")
    ext = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    if ext not in ("csv", "xlsx", "xlsm", "txt", "tsv"):
        raise IntakeError("Upload a .csv or .xlsx file.")
    headers, rows, stats = read_table(content, filename)
    if not rows:
        raise IntakeError("The file has a header row but no data rows.")

    source_key = source if source in SOURCE_LABELS else "other"
    batch = ImportBatch(
        organization_id=ctx.org_id,
        created_by_id=ctx.actor_id, created_by_name=ctx.actor_name,
        display_name=(display_name or list_name or filename or "Import")[:200],
        source_type="xlsx" if ext in ("xlsx", "xlsm") else "csv",
        source_filename=filename,
        status=ImportBatchStatus.MAPPING,
        pipeline=PIPELINE,
        stage="mapping", progress_pct=0,
        acting_user_id=ctx.actor_id, acting_user_name=ctx.actor_name,
        acting_role=ctx.role, acted_as_platform_owner=ctx.acting_as_platform_owner,
        source_label=SOURCE_LABELS[source_key], source_system=source_key,
        source_detail=(source_detail or None),
        source_year=source_year, import_list_name=list_name,
        campaign_purpose=campaign_purpose, offer_hook=offer_hook,
        tags_json=json.dumps([t for t in (tags or []) if t][:50]),
        original_row_count=len(rows),
        total_rows=len(rows),
        headers_json=json.dumps(headers),
        force_new_inquiry=False,
    )
    batch.batch_code = next_batch_code(db, ctx)
    db.add(batch)
    db.flush()
    db.add(ImportBatchFile(
        batch_id=batch.id, organization_id=ctx.org_id, filename=filename,
        content_type="text/csv" if batch.source_type == "csv" else
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        size_bytes=len(content), sha256=hashlib.sha256(content).hexdigest(),
        content=content))
    mapping = F.suggest_mapping(headers, org_custom_field_keys(db, ctx.org_id))
    batch.mapping_json = json.dumps(mapping)
    batch.classification_json = json.dumps(default_classification_config(headers, mapping,
                                                                         rows, db, ctx.org_id))
    batch.update_policy_json = json.dumps(DEFAULT_UPDATE_POLICY)
    batch.analysis_json = json.dumps({"read": {"blank_rows": stats["blank_rows"],
                                               "malformed_rows": stats["malformed_rows"][:200],
                                               "malformed_count": len(stats["malformed_rows"])}})
    return batch


DEFAULT_UPDATE_POLICY = {
    # fill_blanks: an import may only fill fields that are empty on the
    #              existing record. It never replaces a value.
    # overwrite:   fields listed in `fields` may be replaced - except any field
    #              a person corrected by hand, which no import may overwrite.
    "mode": "fill_blanks",
    "fields": [],
}
UPDATABLE_FIELDS = ["first_name", "last_name", "company", "job_title", "email", "phone",
                    "mobile_phone", "street_address", "city", "state", "zip_code", "country",
                    "owner_name", "last_activity_at", "classification", "tags",
                    "custom_fields", "vertical_fields"]


def _column_values(headers, rows, header) -> Dict[str, int]:
    if header not in headers:
        return {}
    i = headers.index(header)
    counts: Dict[str, int] = {}
    for r in rows:
        v = N.clean_text(r[i] if i < len(r) else "")
        if v:
            counts[v] = counts.get(v, 0) + 1
    return counts


def default_classification_config(headers, mapping, rows, db, org_id) -> dict:
    col = next((h for h in headers if mapping.get(h, {}).get("kind") == F.KIND_STANDARD
                and mapping[h].get("target") == "classification"), None)
    cat = org_catalog(db, org_id)
    values = _column_values(headers, rows, col) if col else {}
    vmap = {v: k for v, k in C.suggest_value_map(values, cat).items() if k}
    return {"column": col, "value_map": vmap, "fallback": C.DEFAULT_FALLBACK,
            "date_order": N.DATE_AUTO, "default_country": "US"}


def load_file_rows(db: Session, batch: ImportBatch) -> Tuple[List[str], List[List[str]], dict]:
    f = (db.query(ImportBatchFile)
         .filter(ImportBatchFile.batch_id == batch.id,
                 ImportBatchFile.organization_id == batch.organization_id).first())
    if f is None or f.content is None:
        raise IntakeError("The uploaded file for this batch is no longer available. "
                          "Upload it again.")
    return read_table(bytes(f.content), f.filename or batch.source_filename or "file.csv")


def preview_payload(db: Session, batch: ImportBatch, sample: int = 15) -> dict:
    headers, rows, _ = load_file_rows(db, batch)
    mapping = json.loads(batch.mapping_json or "{}")
    samples = {h: [] for h in headers}
    filled = {h: 0 for h in headers}
    for r in rows:
        for i, h in enumerate(headers):
            v = r[i] if i < len(r) else ""
            if v and v.strip():
                filled[h] += 1
                if len(samples[h]) < 3:
                    samples[h].append(v.strip()[:80])
    return {
        "headers": headers,
        "columns": [{"header": h, "filled": filled[h], "samples": samples[h],
                     "mapping": mapping.get(h) or F.suggest(h)} for h in headers],
        "sample_rows": [dict(zip(headers, r[:len(headers)])) for r in rows[:sample]],
        "row_count": len(rows),
    }


def classification_values(db: Session, batch: ImportBatch) -> dict:
    headers, rows, _ = load_file_rows(db, batch)
    cfg = json.loads(batch.classification_json or "{}")
    col = cfg.get("column")
    cat = org_catalog(db, batch.organization_id)
    values = _column_values(headers, rows, col) if col else {}
    blank_count = len(rows) - sum(values.values())
    vmap = cfg.get("value_map") or {}
    return {
        "column": col,
        "fallback": cfg.get("fallback") or C.DEFAULT_FALLBACK,
        "blank_count": blank_count,
        "values": [{"value": v, "count": n,
                    "classification": vmap.get(v),
                    "suggested": C.suggest_value(v, cat)}
                   for v, n in sorted(values.items(), key=lambda kv: -kv[1])[:500]],
        "catalog": C.payload(cat),
    }


def save_mapping(db: Session, batch: ImportBatch, *, mapping: Optional[dict] = None,
                 classification: Optional[dict] = None,
                 update_policy: Optional[dict] = None) -> List[str]:
    """Validate and store the operator's choices. Returns problems (empty = ok)."""
    headers = json.loads(batch.headers_json or "[]")
    problems: List[str] = []
    if mapping is not None:
        clean = {}
        for h in headers:
            m = mapping.get(h) or {"kind": F.KIND_SOURCE, "target": N.slug(h)}
            kind = m.get("kind")
            tgt = m.get("target")
            if kind in (F.KIND_CUSTOM, F.KIND_VERTICAL, F.KIND_SOURCE) and not tgt:
                tgt = N.slug(h)
            clean[h] = {"kind": kind, "target": tgt}
        problems = F.validate_mapping(headers, clean)
        if problems:
            return problems
        batch.mapping_json = json.dumps(clean)
        cfg = json.loads(batch.classification_json or "{}")
        col = next((h for h in headers if clean[h]["kind"] == F.KIND_STANDARD
                    and clean[h]["target"] == "classification"), None)
        if col != cfg.get("column"):
            _, rows, _ = load_file_rows(db, batch)
            cfg = default_classification_config(headers, clean, rows, db,
                                                batch.organization_id) | {
                "fallback": cfg.get("fallback") or C.DEFAULT_FALLBACK,
                "date_order": cfg.get("date_order") or N.DATE_AUTO}
            batch.classification_json = json.dumps(cfg)
    if classification is not None:
        cfg = json.loads(batch.classification_json or "{}")
        cat = org_catalog(db, batch.organization_id)
        fb = classification.get("fallback", cfg.get("fallback") or C.DEFAULT_FALLBACK)
        if fb not in cat:
            return [f"Unknown fallback classification '{fb}'."]
        vmap = classification.get("value_map", cfg.get("value_map") or {})
        bad = [f"'{v}' -> '{k}'" for v, k in vmap.items() if k and k not in cat]
        if bad:
            return ["Unknown classification in value map: " + ", ".join(bad[:5])]
        order = classification.get("date_order", cfg.get("date_order") or N.DATE_AUTO)
        if order not in (N.DATE_AUTO, N.DATE_MDY, N.DATE_DMY):
            return ["date_order must be auto, mdy or dmy."]
        cfg.update({"fallback": fb, "value_map": {v: k for v, k in vmap.items() if k},
                    "date_order": order})
        batch.classification_json = json.dumps(cfg)
    if update_policy is not None:
        mode = update_policy.get("mode", "fill_blanks")
        if mode not in ("fill_blanks", "overwrite"):
            return ["update_policy.mode must be fill_blanks or overwrite."]
        flds = [f for f in update_policy.get("fields", []) if f in UPDATABLE_FIELDS]
        batch.update_policy_json = json.dumps({"mode": mode, "fields": flds})
    return problems


# ══════════════════════════════════════════════════════════════════════════
# 3. ANALYSIS: PARSE -> NORMALIZE -> MATCH -> CLASSIFY -> OUTREACH -> COUNT
# ══════════════════════════════════════════════════════════════════════════

def _progress(db: Session, batch: ImportBatch, stage: str, pct: int):
    batch.stage = stage
    batch.progress_pct = max(0, min(100, int(pct)))
    batch.heartbeat_at = datetime.utcnow()
    db.commit()


def is_stale(batch: ImportBatch) -> bool:
    if batch.status not in (ImportBatchStatus.PROCESSING, ImportBatchStatus.COMMITTING):
        return False
    hb = batch.heartbeat_at or batch.updated_at
    return bool(hb and datetime.utcnow() - hb.replace(tzinfo=None) > STALE_AFTER)


def _normalize_row(values: Dict[str, str], std: Dict[str, List[str]], cfg: dict,
                   date_orders: Dict[str, str], custom_cols, vertical_cols, source_cols,
                   extra_cells) -> dict:
    """One source row -> one normalized dict (no database access)."""
    from app.services.import_service import _check_email_quality, split_full_name
    from app.services import permission_values as PV

    def g(key):
        for h in std.get(key, []):
            v = values.get(h)
            if not N.blank(v):
                return str(v).strip()
        return None

    reasons: List[dict] = []
    first, last = N.name(g("first_name")), N.name(g("last_name"))
    full = N.name(g("full_name"))
    if not first and not last and full:
        try:
            first, last = split_full_name(full)
            first, last = N.name(first), N.name(last)
        except Exception:  # noqa: BLE001
            pass
    if not full and (first or last):
        full = " ".join(p for p in (first, last) if p)
    company = N.company(g("company"))

    phone_raw = g("phone")
    mobile_raw = g("mobile_phone")
    country = cfg.get("default_country") or "US"
    phone, phone_state = N.phone(phone_raw, country)
    mobile, mobile_state = N.phone(mobile_raw, country)
    if phone_raw and phone_state == N.PHONE_INVALID:
        reasons.append({"code": "invalid_phone", "detail": phone_raw[:40]})
    if mobile_raw and mobile_state == N.PHONE_INVALID:
        reasons.append({"code": "invalid_mobile", "detail": mobile_raw[:40]})
    if mobile and phone == mobile:
        phone = None if not phone_raw else phone
    best_phone = mobile or phone
    best_state = N.PHONE_VALID if best_phone else (
        N.PHONE_INVALID if (phone_raw or mobile_raw) else N.PHONE_MISSING)

    email_raw = g("email")
    email, email_problem = N.email(email_raw)
    quality = _check_email_quality(email) if email else None

    st, st_note = N.state(g("state"))
    zp, zp_note = N.zip_code(g("zip_code"))
    for note in (st_note, zp_note):
        if note:
            reasons.append({"code": note, "detail": ""})
    street = N.address(g("street_address"))
    line2 = N.address(g("address_line2"))
    if street and line2:
        street = f"{street}, {line2}"

    dates = {}
    for key in ("last_activity_date", "source_created_at"):
        raw = g(key)
        if raw:
            hdr = next((h for h in std.get(key, []) if not N.blank(values.get(h))), None)
            dt, prob = N.date(raw, date_orders.get(hdr, N.DATE_AUTO))
            dates[key] = dt
            if prob:
                reasons.append({"code": prob, "detail": f"{key}: {raw[:30]}"})

    # compliance cells, through the platform's ONE permission interpreter.
    # The header decides the polarity ("Do Not Text = Yes" is a denial).
    def consent(key, permission):
        for h in std.get(key, []):
            raw = values.get(h)
            if N.blank(raw):
                continue
            polarity = _polarity(h, permission)
            state, ambiguous = PV.interpret_cell_ex(raw, polarity)
            return PV.to_bool(state), ambiguous, str(raw).strip()
        return None, False, None
    c_email, a_email, r_email = consent("allow_email", PV.EMAIL)
    c_bulk, a_bulk, r_bulk = consent("allow_bulk_email", PV.BULK_EMAIL)
    c_sms, a_sms, r_sms = consent("allow_sms", PV.SMS)
    c_voice, a_voice, r_voice = consent("allow_voice", PV.VOICE)
    if any((a_email, a_bulk, a_sms, a_voice)):
        reasons.append({"code": "ambiguous_consent", "detail": ""})
    dnc_all = N.boolean(g("do_not_contact")) is True
    sms_opt_out = False
    hard_bounce_raw = g("email_hard_bounce")
    hard_bounce = bool(hard_bounce_raw) and N.boolean(hard_bounce_raw) is not False
    unsub_raw = g("email_unsubscribed")
    unsubscribed = N.boolean(unsub_raw) is True
    invalid_flag = N.boolean(g("email_invalid")) is True

    hist_raw = g("historical_customer")
    historical = N.boolean(hist_raw) if hist_raw is not None else None

    src_id = g("source_record_id")
    if src_id and src_id.endswith(".0") and src_id[:-2].isdigit():
        src_id = src_id[:-2]         # 5091051.0 from a spreadsheet is 5091051

    tag_list: List[str] = []
    for h in std.get("tags", []):
        tag_list += [t for t in N.tags(values.get(h)) if t not in tag_list]

    custom = {k: values[h].strip() for h, k in custom_cols if not N.blank(values.get(h))}
    vertical = {k: values[h].strip() for h, k in vertical_cols if not N.blank(values.get(h))}
    source_meta = {k: values[h].strip() for h, k in source_cols if not N.blank(values.get(h))}
    if extra_cells:
        source_meta["_extra_cells"] = extra_cells

    return {
        "first": first, "last": last, "full": full, "company": company,
        "company_key": N.company_key(company), "job_title": N.clean_text(g("job_title")),
        "phone_raw": phone_raw, "mobile_raw": mobile_raw,
        "phone": phone, "mobile": mobile, "best_phone": best_phone, "phone_state": best_state,
        "line_type_raw": g("phone_line_type"), "sms_verification": g("sms_verification"),
        "email_raw": email_raw, "email": email, "email_problem": email_problem,
        "email_quality": quality, "email_verification": g("email_status"),
        "hard_bounce": hard_bounce, "unsubscribed": unsubscribed, "email_invalid": invalid_flag,
        "street": street, "city": N.clean_text(g("city")), "state": st, "zip": zp,
        "country": N.clean_text(g("country")),
        "src_id": src_id, "source": N.clean_text(g("source")),
        "owner_name": N.clean_text(g("owner_name")), "notes": N.clean_text(g("notes")),
        "last_activity_at": dates.get("last_activity_date"),
        "last_activity_raw": g("last_activity_date"),
        "source_created_at": dates.get("source_created_at"),
        "classification_raw": g("classification"),
        "historical_customer": historical,
        "tags": tag_list, "custom": custom, "vertical": vertical, "source_meta": source_meta,
        "consent": {"email": (c_email, r_email), "bulk_email": (c_bulk, r_bulk),
                    "sms": (c_sms, r_sms), "voice": (c_voice, r_voice)},
        "consent_ambiguous": any((a_email, a_bulk, a_sms, a_voice)),
        "dnc_all": dnc_all, "sms_opt_out": sms_opt_out,
        "reasons": reasons,
    }


def _polarity(header: str, permission: str) -> str:
    from app.services import permission_values as PV
    h = PV.normalize_cell(header)
    for col, pol in PV.COLUMN_TABLE.get(permission, ()):
        if PV.normalize_cell(col) == h:
            return pol
    return "grant"


def _mapped_columns(headers, mapping):
    std: Dict[str, List[str]] = {}
    custom, vertical, source = [], [], []
    for h in headers:
        m = mapping.get(h) or {}
        kind, tgt = m.get("kind"), m.get("target")
        if kind == F.KIND_STANDARD and tgt:
            std.setdefault(tgt, []).append(h)
        elif kind == F.KIND_CUSTOM:
            custom.append((h, tgt))
        elif kind == F.KIND_VERTICAL:
            vertical.append((h, tgt))
        elif kind == F.KIND_SOURCE:
            source.append((h, tgt))
    return std, custom, vertical, source


def run_analysis(db: Session, batch_id: str, org_id: str) -> ImportBatch:
    """Stage every row of the batch's file. Idempotent: re-running replaces
    the batch's staged rows. Writes nothing outside import_* tables."""
    batch = (db.query(ImportBatch)
             .filter(ImportBatch.id == batch_id, ImportBatch.organization_id == org_id).first())
    if batch is None:
        raise IntakeError("Batch not found.")
    if batch.pipeline != PIPELINE:
        raise IntakeError("This batch was created by the legacy importer.")
    started = datetime.utcnow()
    try:
        batch.status = ImportBatchStatus.PROCESSING
        batch.error_message = None
        _progress(db, batch, "parsing", 2)
        headers, rows, read_stats = load_file_rows(db, batch)
        mapping = json.loads(batch.mapping_json or "{}")
        cfg = json.loads(batch.classification_json or "{}")
        std, custom_cols, vertical_cols, source_cols = _mapped_columns(headers, mapping)

        # Column-level date order: a column proves its own order when it can.
        date_orders = {}
        for key in ("last_activity_date", "source_created_at"):
            for h in std.get(key, []):
                if cfg.get("date_order") in (N.DATE_MDY, N.DATE_DMY):
                    date_orders[h] = cfg["date_order"]
                else:
                    i = headers.index(h)
                    date_orders[h] = N.slash_date_order_hint(
                        (r[i] if i < len(r) else "") for r in rows) or N.DATE_AUTO

        # ── normalize ──
        _progress(db, batch, "normalizing", 8)
        norm: List[dict] = []
        total = len(rows)
        src_system = batch.source_system or ""
        for n_row, cells in enumerate(rows):
            values = {h: (cells[i] if i < len(cells) else "") for i, h in enumerate(headers)}
            extra = [c for c in cells[len(headers):] if c and c.strip()]
            r = _normalize_row(values, std, cfg, date_orders, custom_cols, vertical_cols,
                               source_cols, extra)
            r["order"] = n_row
            r["row_number"] = n_row + 2
            r["raw"] = {h: v for h, v in values.items() if v not in (None, "")}
            r["src_system"] = src_system if r["src_id"] else ""
            r["first_key"] = N.name_key(r["first"])
            r["last_key"] = N.name_key(r["last"])
            r["addr_key"] = N.address_key(r["street"])
            r["zip5"] = N.zip5(r["zip"])
            r["email"] = r["email"]
            r["phone"] = r["phone"]
            r["mobile"] = r["mobile"]
            norm.append(r)
            if n_row and n_row % 2000 == 0:
                _progress(db, batch, "normalizing", 8 + int(22 * n_row / max(total, 1)))

        # ── within-file duplicates ──
        _progress(db, batch, "deduplicating", 32)
        M.within_batch(norm)

        # ── existing organization records ──
        _progress(db, batch, "matching", 45)
        existing = M.load_existing(db, org_id, norm)
        suppressed = _suppressed_phones(db, org_id, {r["best_phone"] for r in norm
                                                     if r["best_phone"]})
        for r in norm:
            mt, other, keys, notes = M.match_existing(r, existing)
            r["ex_type"], r["ex"], r["ex_keys"], r["ex_notes"] = mt, other, keys, notes

        # ── classify ──
        _progress(db, batch, "classifying", 60)
        cat = org_catalog(db, org_id)
        vmap = cfg.get("value_map") or {}
        fallback = cfg.get("fallback") or C.DEFAULT_FALLBACK
        for r in norm:
            key, how = C.resolve(r["classification_raw"], vmap, fallback, cat)
            cd = cat.get(key) or cat[C.NEEDS_CLASSIFICATION]
            r["classification"], r["classification_source"] = cd.key, how
            r["record_class"], r["creates_lead"] = cd.record_class, cd.creates_lead

        # ── statuses + outreach ──
        _progress(db, batch, "analyzing", 70)
        staged = []
        for r in norm:
            staged.append(_decide(r, batch, suppressed))

        # ── write ──
        _progress(db, batch, "analyzing", 78)
        (db.query(ImportStagedRow)
         .filter(ImportStagedRow.batch_id == batch.id,
                 ImportStagedRow.organization_id == org_id)
         .delete(synchronize_session=False))
        db.commit()
        for i in range(0, len(staged), INSERT_CHUNK):
            db.bulk_insert_mappings(ImportStagedRow, staged[i:i + INSERT_CHUNK])
            db.commit()
            _progress(db, batch, "analyzing", 78 + int(20 * (i + INSERT_CHUNK) / max(len(staged), 1)))

        analysis = compute_analysis(db, batch)
        analysis["read"] = {"blank_rows": read_stats["blank_rows"],
                            "malformed_count": len(read_stats["malformed_rows"]),
                            "malformed_rows": read_stats["malformed_rows"][:200]}
        analysis["timing_seconds"] = round((datetime.utcnow() - started).total_seconds(), 2)
        batch.analysis_json = json.dumps(analysis, default=str)
        batch.total_rows = len(staged)
        batch.new_rows = analysis["match"]["new_records"]
        batch.matched_rows = analysis["match"]["existing_exact"]
        batch.warning_rows = analysis["status"].get(IntakeStatus.NEEDS_REVIEW, 0)
        batch.invalid_rows = analysis["status"].get(IntakeStatus.INVALID, 0)
        batch.pending_rows = len(staged)
        batch.status = ImportBatchStatus.READY_FOR_REVIEW
        _progress(db, batch, "ready_for_review", 100)
        return batch
    except Exception as exc:
        db.rollback()
        batch = db.query(ImportBatch).filter(ImportBatch.id == batch_id).first()
        if batch is not None:
            batch.status = ImportBatchStatus.FAILED
            batch.stage = "failed"
            batch.error_message = (str(exc) if isinstance(exc, IntakeError)
                                   else f"Analysis failed: {type(exc).__name__}: {str(exc)[:300]}")
            db.commit()
        log.exception("intake analysis failed for batch %s", batch_id)
        raise


def _suppressed_phones(db: Session, org_id: str, phones: set) -> set:
    try:
        from app.models.models import SuppressionEntry
    except ImportError:  # pragma: no cover
        return set()
    out = set()
    vals = [p for p in phones if p]
    for i in range(0, len(vals), 900):
        chunk = vals[i:i + 900]
        for (p,) in (db.query(SuppressionEntry.phone)
                     .filter(SuppressionEntry.organization_id == org_id,
                             SuppressionEntry.phone.in_(chunk)).all()):
            out.add(p)
    return out


_LEGACY_DUP = {
    IntakeStatus.DUPLICATE: ImportDuplicateStatus.WITHIN_BATCH_DUPLICATE,
    IntakeStatus.BLOCKED: ImportDuplicateStatus.DNC_BLOCKED,
}


def _decide(r: dict, batch: ImportBatch, suppressed: set) -> dict:
    """Turn one normalized, matched, classified row into a staged-row mapping."""
    reasons = list(r["reasons"])
    ex, ex_type = r.get("ex"), r.get("ex_type")
    dnc_existing = bool(ex and ex.get("dnc"))
    existing_email_denied = bool(ex and ex.get("allow_email") is False)

    # ── OUTREACH STATUS ──
    line = E.line_type(r["line_type_raw"], r["sms_verification"], bool(r["mobile"]))
    if not r["best_phone"]:
        line = None
    c_sms = r["consent"]["sms"][0]
    sms, sms_reasons = E.sms_status(
        phone_state=r["phone_state"], line=line or E.LINE_UNKNOWN,
        dnc=r["dnc_all"] or dnc_existing, suppressed=r["best_phone"] in suppressed,
        opted_out=r["sms_opt_out"], consent_sms=c_sms,
        sms_verification=r["sms_verification"])
    email, email_reasons = E.email_status(
        email_norm=r["email"], email_raw_present=bool(r["email_raw"]),
        format_problem=r["email_problem"], quality_problem=r["email_quality"],
        verification=r["email_verification"], hard_bounce=r["hard_bounce"],
        unsubscribed=r["unsubscribed"], marked_invalid=r["email_invalid"],
        suppressed=r["dnc_all"] or dnc_existing or existing_email_denied,
        consent_email=r["consent"]["email"][0])

    usable_phone = bool(r["best_phone"]) and sms not in (SmsStatus.INVALID, SmsStatus.DNC)
    usable_email = email in E.USABLE_EMAIL_STATUSES
    needs_enrichment = not usable_phone and not usable_email

    # ── IMPORT STATUS (first matching rule wins) ──
    has_identity = any((r["first"], r["last"], r["full"], r["company"]))
    has_anything = has_identity or r["best_phone"] or r["email"] or r["src_id"] \
        or r["phone_raw"] or r["email_raw"]
    resolution = None
    if not has_anything:
        status = IntakeStatus.INVALID
        reasons.append({"code": "no_identity", "detail": "no name, company, phone, email or id"})
    elif r["dup_type"] == MatchType.EXACT:
        status = IntakeStatus.DUPLICATE
        reasons.append({"code": "duplicate_in_file",
                        "detail": f"same {'/'.join(r['dup_keys'])} as row {r['dup_of'] + 2}"})
        # Lossless by default: the row's values fill blanks on the first
        # occurrence and its source id is kept as an alternate id.
        resolution = DuplicateResolution.MERGE
    elif dnc_existing:
        status = IntakeStatus.BLOCKED
        reasons.append({"code": "existing_record_dnc", "detail": ex.get("id")})
    elif ex_type == MatchType.EXACT:
        status = IntakeStatus.EXISTING_MATCH
        resolution = DuplicateResolution.UPDATE_EXISTING
        reasons.append({"code": "existing_record_match", "detail": "/".join(r["ex_keys"])})
    else:
        status = None
        if ex_type == MatchType.POSSIBLE:
            reasons.append({"code": "possible_existing_match", "detail": "/".join(r["ex_keys"])})
            resolution = DuplicateResolution.REVIEW
        if r["dup_type"] == MatchType.POSSIBLE:
            reasons.append({"code": "possible_duplicate_in_file",
                            "detail": f"{'/'.join(r['dup_keys'])} ~ row {r['dup_of'] + 2}"})
            resolution = DuplicateResolution.REVIEW
        if r["classification"] == C.NEEDS_CLASSIFICATION:
            reasons.append({"code": "unrecognized_classification",
                            "detail": (r["classification_raw"] or "")[:60]})
        review_codes = {"possible_existing_match", "possible_duplicate_in_file",
                        "unrecognized_classification", "ambiguous_consent", "ambiguous_date"}
        if any(x["code"] in review_codes for x in reasons):
            status = IntakeStatus.NEEDS_REVIEW
        elif needs_enrichment:
            status = IntakeStatus.NEEDS_ENRICHMENT
        else:
            status = IntakeStatus.READY
    for n in (r.get("dup_notes") or []) + (r.get("ex_notes") or []):
        reasons.append({"code": n, "detail": ""})
    if needs_enrichment and status != IntakeStatus.NEEDS_ENRICHMENT:
        reasons.append({"code": "no_usable_channel", "detail": ""})

    historical = r["historical_customer"]
    matched_customer = bool(ex and ex.get("record_class") in ("customer", "previous_customer"))

    legacy_dup = _LEGACY_DUP.get(status) or {
        MatchType.EXACT: ImportDuplicateStatus.MATCHED_EXISTING,
        MatchType.POSSIBLE: ImportDuplicateStatus.POSSIBLE_DUPLICATE,
    }.get(ex_type, ImportDuplicateStatus.NEW)
    matched_lead_id = ex.get("lead_id") if ex and ex_type != MatchType.NEW else None
    matched_contact_id = (ex.get("id") if ex and ex.get("kind") == "org_contact"
                          and ex_type != MatchType.NEW else None)
    consent = r["consent"]

    return {
        "batch_id": batch.id, "organization_id": batch.organization_id,
        "row_number": r["row_number"],
        "raw_data": json.dumps(r["raw"], default=str),
        "first_name": r["first"], "last_name": r["last"], "full_name": r["full"],
        "company": r["company"], "company_norm": r["company_key"],
        "phone_raw": r["phone_raw"], "phone_normalized": r["best_phone"],
        "mobile_phone_raw": r["mobile_raw"], "mobile_phone_normalized": r["mobile"],
        "phone_type": ("known_mobile" if line == E.LINE_MOBILE else
                       "known_landline" if line == E.LINE_LANDLINE else
                       ("unknown" if r["best_phone"] else None)),
        "phone_line_type": line,
        "email_raw": r["email_raw"], "email_normalized": r["email"],
        "street_address": r["street"], "city": r["city"], "state": r["state"],
        "zip_code": r["zip"],
        "source_category": (r["source"] or None),
        "source_id": r["src_id"], "source_id_type": (batch.source_system + "_record_id")
        if r["src_id"] else None,
        "source_system": r["src_system"] or None, "source_record_id": r["src_id"],
        "last_activity_date": r["last_activity_at"], "last_activity_date_raw": r["last_activity_raw"],
        "consent_email": consent["email"][0], "consent_email_raw": consent["email"][1],
        "consent_bulk_email": consent["bulk_email"][0],
        "consent_bulk_email_raw": consent["bulk_email"][1],
        "consent_sms": consent["sms"][0], "consent_sms_raw": consent["sms"][1],
        "consent_voice": consent["voice"][0], "consent_voice_raw": consent["voice"][1],
        "consent_review_required": bool(r["consent_ambiguous"]),
        "validation_status": (ImportValidationStatus.INVALID if status == IntakeStatus.INVALID
                              else ImportValidationStatus.WARNING if reasons
                              else ImportValidationStatus.VALID),
        "validation_errors": json.dumps([x["code"] for x in reasons]) if reasons else None,
        "duplicate_status": legacy_dup,
        "match_confidence": (ImportMatchConfidence.HIGH if ex_type == MatchType.EXACT else
                             ImportMatchConfidence.MEDIUM if ex_type == MatchType.POSSIBLE
                             else ImportMatchConfidence.NONE),
        "matched_lead_id": matched_lead_id,
        "review_status": ImportRowReviewStatus.PENDING,
        # universal columns
        "intake_status": status,
        "status_reasons": json.dumps(reasons),
        "record_class": r["record_class"], "classification": r["classification"],
        "classification_source": r["classification_source"],
        "classification_raw": (r["classification_raw"] or None),
        "creates_lead": bool(r["creates_lead"]),
        "historical_customer": historical if historical is not None else (
            True if matched_customer else None),
        "needs_enrichment": needs_enrichment,
        "sms_status": sms, "email_status": email,
        "email_status_raw": r["email_verification"],
        "match_type": (MatchType.EXACT if status == IntakeStatus.DUPLICATE else ex_type),
        "match_target_type": ("staged_row" if status == IntakeStatus.DUPLICATE
                              else (ex.get("kind") if ex and ex_type != MatchType.NEW else None)),
        "matched_contact_id": matched_contact_id,
        "match_keys": json.dumps(r["dup_keys"] if status == IntakeStatus.DUPLICATE
                                 else r["ex_keys"]),
        "duplicate_resolution": resolution,
        "normalized_json": json.dumps({
            "phone_e164": r["phone"],
            "job_title": r["job_title"], "country": r["country"], "owner_name": r["owner_name"],
            "notes": r["notes"], "source": r["source"],
            "source_created_at": r["source_created_at"], "dup_of_row": (
                r["dup_of"] + 2 if r["dup_of"] is not None else None),
            "outreach_reasons": sms_reasons + email_reasons,
            "matched_record_class": ex.get("record_class") if ex else None,
            "matched_customer": matched_customer,
            "source_meta": r["source_meta"],
        }, default=str),
        "custom_fields_json": json.dumps(r["custom"]) if r["custom"] else None,
        "vertical_fields_json": json.dumps(r["vertical"]) if r["vertical"] else None,
        "tags_json": json.dumps(r["tags"]) if r["tags"] else None,
    }


# ══════════════════════════════════════════════════════════════════════════
# 4. THE PREVIEW COUNTS
# ══════════════════════════════════════════════════════════════════════════

def compute_analysis(db: Session, batch: ImportBatch) -> dict:
    """Every number on the Analyze screen, counted from the staged rows the
    commit will read - so the preview and the commit cannot disagree."""
    from sqlalchemy import func
    q = db.query(ImportStagedRow).filter(ImportStagedRow.batch_id == batch.id,
                                         ImportStagedRow.organization_id == batch.organization_id)
    total = q.count()

    def grouped(col):
        return {k or "none": n for k, n in
                q.with_entities(col, func.count()).group_by(col).all()}

    status = grouped(ImportStagedRow.intake_status)
    sms = grouped(ImportStagedRow.sms_status)
    email = grouped(ImportStagedRow.email_status)
    line = grouped(ImportStagedRow.phone_line_type)
    cls = grouped(ImportStagedRow.classification)
    rclass = grouped(ImportStagedRow.record_class)
    match = grouped(ImportStagedRow.match_type)

    usable_email_q = ImportStagedRow.email_status.in_(E.USABLE_EMAIL_STATUSES)
    usable_phone_q = (ImportStagedRow.phone_normalized.isnot(None)
                      & ~ImportStagedRow.sms_status.in_((SmsStatus.INVALID, SmsStatus.DNC)))
    phone_email = q.filter(usable_phone_q, usable_email_q).count()
    phone_only = q.filter(usable_phone_q, ~usable_email_q).count()
    email_only = q.filter(~usable_phone_q, usable_email_q).count()
    none_direct = total - phone_email - phone_only - email_only

    dup_in_file = status.get(IntakeStatus.DUPLICATE, 0)
    possible_in_file = q.filter(ImportStagedRow.status_reasons.like(
        '%possible_duplicate_in_file%')).count()
    existing_exact = q.filter(ImportStagedRow.match_type == MatchType.EXACT,
                              ImportStagedRow.intake_status != IntakeStatus.DUPLICATE).count()
    existing_possible = q.filter(ImportStagedRow.match_type == MatchType.POSSIBLE).count()
    exact_customer = q.filter(ImportStagedRow.match_type == MatchType.EXACT,
                              ImportStagedRow.intake_status != IntakeStatus.DUPLICATE,
                              ImportStagedRow.normalized_json.like('%"matched_customer": true%')
                              ).count()
    possible_customer = q.filter(ImportStagedRow.match_type == MatchType.POSSIBLE,
                                 ImportStagedRow.normalized_json.like(
                                     '%"matched_customer": true%')).count()
    historical = q.filter(ImportStagedRow.historical_customer.is_(True)).count()
    needs_enrichment = q.filter(ImportStagedRow.needs_enrichment.is_(True)).count()
    valid_phones = q.filter(ImportStagedRow.phone_normalized.isnot(None)).count()
    invalid_phones = sms.get(SmsStatus.INVALID, 0)
    would_lead = q.filter(ImportStagedRow.creates_lead.is_(True),
                          ImportStagedRow.intake_status.in_((IntakeStatus.READY,)),
                          ImportStagedRow.needs_enrichment.isnot(True)).count()

    return {
        "total_rows": total,
        "unique_records": total - dup_in_file,
        "status": status,
        "match": {
            "duplicates_in_file": dup_in_file,
            "possible_duplicates_in_file": possible_in_file,
            "existing_exact": existing_exact,
            "existing_possible": existing_possible,
            "new_records": match.get(MatchType.NEW, 0) + match.get("none", 0),
            "exact_customer_matches": exact_customer,
            "possible_customer_matches": possible_customer,
        },
        "phones": {
            "valid_format": valid_phones,
            "missing": sms.get(SmsStatus.NO_PHONE, 0),
            "invalid": invalid_phones,
            "line_type": {"mobile": line.get("mobile", 0), "landline": line.get("landline", 0),
                          "voip": line.get("voip", 0), "unknown": line.get("unknown", 0)},
        },
        "emails": {
            "usable": q.filter(usable_email_q).count(),
            "ready": email.get(EmailStatus.READY, 0),
            "review": email.get(EmailStatus.REVIEW, 0),
            "invalid": email.get(EmailStatus.INVALID, 0),
            "hard_bounce": email.get(EmailStatus.HARD_BOUNCE, 0),
            "unsubscribed": email.get(EmailStatus.UNSUBSCRIBED, 0),
            "suppressed": email.get(EmailStatus.SUPPRESSED, 0),
            "missing": email.get(EmailStatus.NO_EMAIL, 0),
        },
        "coverage": {"phone_and_email": phone_email, "phone_only": phone_only,
                     "email_only": email_only, "no_direct_contact": none_direct},
        "sms": sms,
        "email_status": email,
        "classification": cls,
        "record_class": rclass,
        "historical_customers": historical,
        "needs_enrichment": needs_enrichment,
        "ready": status.get(IntakeStatus.READY, 0),
        "review": status.get(IntakeStatus.NEEDS_REVIEW, 0),
        "blocked": status.get(IntakeStatus.BLOCKED, 0),
        "invalid": status.get(IntakeStatus.INVALID, 0),
        "failed": status.get(IntakeStatus.FAILED, 0),
        "would_activate_as_leads_if_ready_committed": would_lead,
        "sms_ready": sms.get(SmsStatus.READY, 0),
        "email_ready": email.get(EmailStatus.READY, 0),
    }


# ══════════════════════════════════════════════════════════════════════════
# 5. CATEGORY DRILL-DOWN
# ══════════════════════════════════════════════════════════════════════════

def category_filter(q, category: Optional[str]):
    """Apply one Analyze-screen category to a staged-row query."""
    if not category or category == "all":
        return q
    kind, _, val = category.partition(":")
    R = ImportStagedRow
    usable_email = R.email_status.in_(E.USABLE_EMAIL_STATUSES)
    usable_phone = R.phone_normalized.isnot(None) & ~R.sms_status.in_(
        (SmsStatus.INVALID, SmsStatus.DNC))
    if kind == "status":
        return q.filter(R.intake_status == val)
    if kind == "class":
        return q.filter(R.classification == val)
    if kind == "record_class":
        return q.filter(R.record_class == val)
    if kind == "sms":
        return q.filter(R.sms_status == val)
    if kind == "email":
        return q.filter(R.email_status == val)
    if kind == "line":
        return q.filter(R.phone_line_type == val)
    if kind == "coverage":
        return q.filter({"phone_and_email": usable_phone & usable_email,
                         "phone_only": usable_phone & ~usable_email,
                         "email_only": ~usable_phone & usable_email,
                         "no_direct_contact": ~usable_phone & ~usable_email}.get(
            val, R.id.is_(None)))
    if kind == "match":
        if val == "existing_exact":
            return q.filter(R.match_type == MatchType.EXACT,
                            R.intake_status != IntakeStatus.DUPLICATE)
        if val == "existing_possible":
            return q.filter(R.match_type == MatchType.POSSIBLE)
        if val == "duplicates_in_file":
            return q.filter(R.intake_status == IntakeStatus.DUPLICATE)
        if val == "possible_duplicates_in_file":
            return q.filter(R.status_reasons.like("%possible_duplicate_in_file%"))
        if val == "new":
            return q.filter((R.match_type == MatchType.NEW) | R.match_type.is_(None))
        if val == "customer_exact":
            return q.filter(R.match_type == MatchType.EXACT,
                            R.normalized_json.like('%"matched_customer": true%'))
    if kind == "flag":
        if val == "historical_customer":
            return q.filter(R.historical_customer.is_(True))
        if val == "needs_enrichment":
            return q.filter(R.needs_enrichment.is_(True))
        if val == "creates_lead":
            return q.filter(R.creates_lead.is_(True))
    if kind == "phones":
        if val == "valid":
            return q.filter(R.phone_normalized.isnot(None))
        if val == "missing":
            return q.filter(R.sms_status == SmsStatus.NO_PHONE)
    if kind == "emails" and val == "usable":
        return q.filter(usable_email)
    if kind == "reason":
        return q.filter(R.status_reasons.like(f'%"{val}"%'))
    raise IntakeError(f"Unknown category '{category}'.")


def row_payload(r: ImportStagedRow) -> dict:
    def j(s, d):
        try:
            return json.loads(s) if s else d
        except Exception:  # noqa: BLE001
            return d
    norm = j(r.normalized_json, {})
    return {
        "id": r.id, "row_number": r.row_number,
        "first_name": r.first_name, "last_name": r.last_name, "full_name": r.full_name,
        "company": r.company, "email": r.email_normalized, "email_raw": r.email_raw,
        "phone": r.phone_normalized, "phone_raw": r.phone_raw,
        "mobile": r.mobile_phone_normalized, "street_address": r.street_address,
        "city": r.city, "state": r.state, "zip_code": r.zip_code,
        "source_record_id": r.source_record_id,
        "intake_status": r.intake_status, "reasons": j(r.status_reasons, []),
        "classification": r.classification, "classification_source": r.classification_source,
        "classification_raw": r.classification_raw, "record_class": r.record_class,
        "creates_lead": r.creates_lead, "historical_customer": r.historical_customer,
        "needs_enrichment": r.needs_enrichment,
        "sms_status": r.sms_status, "email_status": r.email_status,
        "phone_line_type": r.phone_line_type,
        "match_type": r.match_type, "match_target_type": r.match_target_type,
        "matched_contact_id": r.matched_contact_id, "matched_lead_id": r.matched_lead_id,
        "match_keys": j(r.match_keys, []), "duplicate_resolution": r.duplicate_resolution,
        "dup_of_row": norm.get("dup_of_row"),
        "outreach_reasons": norm.get("outreach_reasons", []),
        "custom_fields": j(r.custom_fields_json, {}),
        "vertical_fields": j(r.vertical_fields_json, {}),
        "tags": j(r.tags_json, []),
        "commit_action": r.commit_action, "committed_contact_id": r.committed_contact_id,
        "committed_lead_id": r.committed_lead_id,
    }
